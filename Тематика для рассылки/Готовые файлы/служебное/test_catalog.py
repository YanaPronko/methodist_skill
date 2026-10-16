"""Integration checks for the simplified survey and the refreshed OKBO source."""
from pathlib import Path
import shutil
import tempfile
import unittest

from openpyxl import load_workbook
import catalog_tools as m


class CatalogueTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data=m.catalog();cls.template=m.ROOT/m.SURVEY
        cls.shared=next(t for t in cls.data['topics'] if len(m.codes(t))>=3)
        cls.single=next(t for t in cls.data['topics'] if len(m.codes(t))==1)
        cls.other=next(t for t in cls.data['topics'] if t['id'] not in [cls.shared['id'],cls.single['id']])
        cls.row_by_id={t['id']:i for i,t in enumerate(cls.data['topics'],m.FIRST_TOPIC_ROW)}

    def setUp(self):self.temp=tempfile.TemporaryDirectory();self.folder=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()

    def response(self,filename,name,answers):
        w=load_workbook(self.template);s=w['Выбор тем'];s['B4']=name;s['B5']='Выбор завершён'
        for topic,ready,comment in answers:
            row=self.row_by_id[topic['id']];s.cell(row,5).value=ready;s.cell(row,6).value=comment
        path=self.folder/filename;w.save(path);w.close();return path

    def test_simplified_layout_and_links(self):
        w=load_workbook(self.template);s=w['Выбор тем']
        self.assertEqual(w.sheetnames,['Выбор тем','Содержание тем','Инструкция'])
        self.assertEqual(s.max_column,6)
        self.assertEqual([s.cell(m.HEADER_ROW,i).value for i in range(1,7)],m.HEADERS)
        self.assertNotIn('LecturerContact',w.defined_names)
        self.assertIsNone(s['B4'].value)
        self.assertFalse(any(c.data_type=='f' for row in s for c in row))
        text=' '.join(str(c.value or '') for row in s for c in row)
        for removed in ['Контакт','Выбрано тем','Для каких программ','Количество тем']:
            self.assertNotIn(removed,text)
        for row,t in enumerate(self.data['topics'],m.FIRST_TOPIC_ROW):
            self.assertEqual(s.cell(row,1).value,t['id'])
            self.assertIsNone(s.cell(row,5).value);self.assertIsNone(s.cell(row,6).value)
            target=s.cell(row,4).hyperlink.location.split('!')[1]
            self.assertEqual(w['Содержание тем'][target].value,t['title'])
        self.assertFalse(any(c.comment for sh in w for row in sh for c in row))
        self.assertFalse(w._external_links)
        for sh in w:
            for row in sh:
                for c in row:
                    if c.hyperlink and c.hyperlink.location:
                        self.assertNotIn("'Программы'!",c.hyperlink.location)
        w.close()

    def test_updated_okbo_coverage(self):
        d=self.data;update=d['latest_source_update'];by_id={t['id']:t for t in d['topics']}
        self.assertEqual(len(d['programs']),18)
        self.assertEqual(d['stats']['utp_rows'],327)
        self.assertEqual(d['stats']['program_only_rows'],0)
        self.assertEqual(len(update['coverage']),32)
        program=next(p for p in d['programs'] if p['code']=='ОКБО')
        self.assertEqual(program['program_sha256'],update['source']['sha256'])
        self.assertEqual(program['utp_sha256'],update['utp_source']['sha256'])
        self.assertEqual(program['utp_url'],'https://drive.google.com/file/d/1bOUvJIn8iJR-cZNs_B-_NVS1Z9rSGipG/view')
        self.assertNotEqual(program['utp_sha256'],update['previous_utp_source']['sha256'])
        self.assertEqual(len(by_id),len(d['topics']))
        okbo=[u for t in d['topics'] for u in t['usages'] if u['program']=='ОКБО']
        self.assertEqual(len(okbo),32)
        self.assertTrue(all(u['content'] and u['utp_table'] and u['utp_row'] for u in okbo))
        self.assertEqual(len({h for u in okbo for h in u['source_headings']}),32)
        for row in update['coverage']:
            self.assertIn(row['id'],by_id)
            uses=[u for u in by_id[row['id']]['usages'] if u['program']=='ОКБО']
            self.assertEqual(len(uses),1)
            self.assertEqual((uses[0]['utp_title'],uses[0]['utp_table'],uses[0]['utp_row']),
                             (row['utp_title'],row['utp_table'],row['utp_row']))
        self.assertEqual({t['title'] for t in update['former_program_only_topics_now_in_utp']},{
            'Подключение различных источников событий',
            'Нормализация событий, корреляция, обработка ложных событий',
            'Совместное функционирование систем SIEM, SOAR и TIP'})
        for item in update['former_program_only_topics_now_in_utp']:
            uses=[u for u in by_id[item['id']]['usages'] if u['program']=='ОКБО']
            self.assertTrue(all(u['source_kind']=='utp' and u['utp_row'] and u['content'] for u in uses))
        for item in update['removed_obsolete_topics']:
            self.assertNotIn(item['id'],by_id)
        self.assertEqual(update['remaining_missing_descriptions'],[])
        self.assertEqual(sum(not u['content'] for t in d['topics'] for u in t['usages']),0)

    def test_topic_readiness_and_partial_comments(self):
        a=self.response('a.xlsx','Тестовый лектор 1',[(self.shared,'Да',''),(self.single,'Частично','Только теория')])
        b=self.response('b.xlsx','Тестовый лектор 2',[(self.shared,'Да',''),(self.other,'Да','')])
        result=m.make_registry(self.data,[a,b],self.folder/'svod.xlsx')
        self.assertEqual(result['forms'],2);self.assertEqual(result['marks'],4)
        expected=2*len(m.codes(self.shared))+len(m.codes(self.single))+len(m.codes(self.other))
        self.assertEqual(result['topic_program_lecturer_links'],expected)
        w=load_workbook(self.folder/'svod.xlsx',data_only=True)
        self.assertNotIn('Программы',w.sheetnames)
        self.assertNotIn('Контакт / код',[c.value for c in w['Реестр'][5]])
        rows=list(w['Реестр'].iter_rows(min_row=6,values_only=True))
        for name in ['Тестовый лектор 1','Тестовый лектор 2']:
            shared=[r for r in rows if r[2]==name and r[7]==self.shared['id']]
            self.assertEqual({r[1] for r in shared},set(m.codes(self.shared)))
        partial=[r for r in rows if r[2]=='Тестовый лектор 1' and r[3]=='Частично']
        self.assertTrue(partial);self.assertTrue(all(r[6]=='Уточнить объём' and r[4]=='Только теория' for r in partial))
        self.assertTrue(any(not r[2] and r[6]=='Нет отметок' for r in rows));w.close()

    def test_sorting_preserves_answers(self):
        p=self.response('a.xlsx','Тест',[(self.shared,'Да',''),(self.single,'Частично','Ограничение по программе')])
        original=m.read_response(p,self.data);self.assertNotIn('contact',original)
        w=load_workbook(p);s=w['Выбор тем'];last=m.HEADER_ROW+len(self.data['topics'])
        values=[[s.cell(r,c).value for c in range(1,7)] for r in range(m.FIRST_TOPIC_ROW,last+1)]
        for r,row in zip(range(m.FIRST_TOPIC_ROW,last+1),reversed(values)):
            for c,value in enumerate(row,1):s.cell(r,c).value=value
        w.save(p);w.close()
        self.assertEqual(m.read_response(p,self.data)['answers'],original['answers'])

    def test_duplicate_and_conflicting_versions(self):
        a=self.response('a.xlsx','Тест',[(self.shared,'Да','')])
        b=self.folder/'duplicate.xlsx';shutil.copyfile(a,b)
        people,duplicates=m.responses(self.data,[a,b])
        self.assertEqual(len(people),1);self.assertEqual(len(duplicates),1)
        w=load_workbook(b);w['Выбор тем'].cell(self.row_by_id[self.shared['id']],5).value='Частично';w.save(b);w.close()
        with self.assertRaises(ValueError):m.make_registry(self.data,[a,b],self.folder/'bad.xlsx')
        self.assertFalse((self.folder/'bad.xlsx').exists())

    def test_invalid_responses_and_old_versions(self):
        p=self.response('a.xlsx','Тест',[(self.shared,'Да','')]);row=self.row_by_id[self.shared['id']]
        for cell,value in [('B4',''),('F5','ROZUM-2026-10-15-v1'),('F5','ROZUM-2026-10-16-v2'),(f'E{row}','Неверно'),(f'F{row}','=1+1'),(f'A{row}','T-UNKNOWN')]:
            w=load_workbook(p);w['Выбор тем'][cell]=value;q=self.folder/'bad.xlsx';w.save(q);w.close()
            with self.assertRaises(ValueError,msg=cell):m.read_response(q,self.data)

    def test_empty_response_and_distinct_full_names(self):
        a=self.response('a.xlsx','Одинаковое ФИО (организация А)',[])
        b=self.response('b.xlsx','Одинаковое ФИО (организация Б)',[])
        people,_=m.responses(self.data,[a,b]);self.assertEqual(len(people),2)
        result=m.make_registry(self.data,[a],self.folder/'empty.xlsx');self.assertEqual(result['marks'],0)
        pairs={(t['id'],u['program']) for t in self.data['topics'] for u in t['usages']}
        w=load_workbook(self.folder/'empty.xlsx');self.assertEqual(w['Реестр'].max_row-5,len(pairs));w.close()


if __name__=='__main__':unittest.main()
