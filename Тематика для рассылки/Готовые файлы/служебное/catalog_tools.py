"""Offline lecturer survey: canonical topics, program-specific scope and merging."""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
import argparse
import json
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

import brand as b

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ['Код темы', 'Тема', 'В каких программах', 'Содержание', 'Могу читать', 'Комментарий преподавателя']
HEADER_ROW = 9
FIRST_TOPIC_ROW = HEADER_ROW + 1
SURVEY = '01_Анкета_преподавателя_Rozum.xlsx'
REGISTRY = '02_Реестр_лекторов.xlsx'
AUDIT = '03_Проверка_источников.xlsx'


def catalog():
    return json.loads((ROOT / 'служебное/каталог.json').read_text(encoding='utf-8'))


def norm(value):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', str(value or ''))).strip().casefold()


def codes(topic):
    return sorted({u['program'] for u in topic['usages']})


def new_book():
    wb = Workbook()
    wb.remove(wb.active)
    b.brand_workbook(wb)
    wb.properties.creator = 'Rozum'
    return wb


def sheet(wb, name, widths, title=None, stripe=4):
    s = wb.create_sheet(name)
    b.base_sheet(s, widths)
    b.brand_title(s, len(widths), title or name.upper())
    s.oddFooter.center.text = 'rozum · &P / &N'
    if stripe:
        b.brand_stripe(s, stripe, len(widths))
    return s


def named(wb, name, cell):
    wb.defined_names.add(DefinedName(name, attr_text=f"'Выбор тем'!${cell[0]}${cell[1:]}"))


def dropdown(s, target, values, strict=True):
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
    dv.showDropDown = False
    dv.showInputMessage = False
    dv.showErrorMessage = strict
    dv.errorStyle = 'stop'
    dv.errorTitle = 'Выберите значение из списка'
    dv.error = 'Выберите одно из предложенных значений либо очистите ячейку.'
    s.add_data_validation(dv)
    dv.add(target)


def source_link(cell, label, url):
    b.text(cell, label)
    cell.hyperlink = url
    cell.font = Font(name='Calibri', size=11, color=b.TEAL, underline='single')
    cell.alignment = Alignment(wrap_text=True, vertical='center')


def chunks(text, limit=850):
    for para in text.splitlines():
        para = para.strip()
        if not para:
            continue
        while len(para) > limit:
            split = para.rfind(' ', 0, limit)
            if split < 100:
                split = limit
            yield para[:split]
            para = para[split:].lstrip()
        if para:
            yield para


def add_details(wb, data, back_sheet, back_col, first_row, id_col):
    s = sheet(wb, 'Содержание тем', [15, 29, 29, 29, 29, 29, 31])
    b.band(s, 2, 1, 7, 'Темы из УТП и программ. Ниже — фактические тексты источников; различающиеся варианты сохранены отдельно.', height=29)
    s.merge_cells('A3:C3')
    b.link(s['A3'], '← К перечню тем', back_sheet, 1)
    s.row_dimensions[3].height=26
    anchors = {}
    row = 5
    program_by_code = {p['code']:p for p in data['programs']}
    for t in data['topics']:
        anchors[t['id']] = row
        b.band(s, row, 1, 7, t['title'], color=b.TEAL, size=13, bold=True,
               height=max(38, 19*(len(t['title'])//120+1)+10)); row += 1
        b.band(s, row, 1, 7, t['id']+' · Программы: '+', '.join(codes(t)), height=26); row += 1
        variants = defaultdict(list)
        for u in t['usages']:
            variants[(u['content'],u['status'],u['note'])].append(u)
        for (content,status,note), uses in variants.items():
            b.band(s, row, 1, 7, 'Вариант для: '+', '.join(sorted({u['program'] for u in uses})),
                   color=b.INK, bold=True, height=27); row += 1
            if note:
                b.band(s, row, 1, 7, note, color=b.NOTICE, height=44 if len(note)>170 else 31); row += 1
            if not content:
                content='Описание этой темы не найдено в предоставленной программе. В УТП тема присутствует. Текст не дополнен предположениями; требуется уточнение у организатора.'
            for part in chunks(content):
                b.band(s, row, 1, 7, part, color=b.WHITE, size=11,
                       height=max(30, 16*(len(part)//145+2))); row += 1
            for u in uses:
                p = program_by_code[u['program']]
                if u.get('utp_table') is not None:
                    location = f"УТП: таблица {u['utp_table']}, строка {u['utp_row']}"
                else:
                    location = 'Тема добавлена непосредственно из обновлённой программы'
                locator = f"{u['program']} — {p['title']}\n{u['section']} · {location}"
                b.band(s, row, 1, 7, locator, height=51); row += 1
                if u['source_headings']:
                    b.band(s, row, 1, 7, 'В программе: '+' / '.join(u['source_headings']),
                           height=max(30,17*(sum(map(len,u['source_headings']))//155+1)+12)); row += 1
                s.merge_cells(start_row=row,start_column=1,end_row=row,end_column=3)
                s.merge_cells(start_row=row,start_column=4,end_row=row,end_column=7)
                source_link(s.cell(row,1), 'Исходная программа '+u['program']+' ↗', p['program_url'])
                source_link(s.cell(row,4), ('УТП для сверки ' if u.get('utp_table') is None else 'Исходный УТП ')+u['program']+' ↗', p['utp_url'])
                s.row_dimensions[row].height=25;row+=1
        s.merge_cells(start_row=row,start_column=1,end_row=row,end_column=7)
        c=s.cell(row,1)
        c.value=f'=HYPERLINK("#\'{back_sheet}\'!{back_col}"&MATCH("{t["id"]}",\'{back_sheet}\'!${id_col}:${id_col},0),"← Вернуться к теме")'
        c.font=Font(name='Calibri',size=11,color=b.TEAL,underline='single',bold=True)
        s.row_dimensions[row].height=27;row+=2
    s.freeze_panes='A5'
    s.print_title_rows='1:3'
    s.print_area=f'A1:G{row-1}'
    return anchors


def add_instructions(wb, data, admin=False):
    s=sheet(wb,'Инструкция',[25,30,30,30,30,30,28],stripe=2)
    notes = [
        ('1 · Сохраните свою копию', 'Откройте файл в Excel или LibreOffice Calc. Предварительный просмотр в чате может не поддерживать редактирование.'),
        ('2 · Представьтесь', 'Укажите ФИО полностью. Дату можно не заполнять.'),
        ('3 · Посмотрите содержание', 'Нажмите «Открыть →» напротив темы. Если тема встречается в нескольких программах, изучите соответствующие варианты. Интернет для содержания не нужен.'),
        ('4 · Отметьте готовность', '«Да» — готовы читать тему в приведённом объёме. «Частично» — есть ограничения по содержанию или программе: поясните их в комментарии. Пустая отметка означает «не указано», а не отказ.'),
        ('5 · Завершите и отправьте', 'Выберите «Выбор завершён», сохраните свою копию как «Анкета_Фамилия.xlsx» и отправьте организатору лично. В общий чат предназначен только пустой шаблон.'),
        ('Что находится в каталоге', 'Темы из программ и УТП. Повторяющиеся названия объединены; различия содержания сохранены по программам. Полные названия программ и ссылки на источники находятся в карточках тем.'),
        ('Расхождения источников', 'Темы без отдельного описания и несовпадения программы с УТП помечены в карточках. В ПСАЭ часть тем входит в общий практикум. Подробности — в отчёте проверки источников у организатора.'),
    ]
    if admin:
        notes=[('Для организатора','Этот файл содержит свод всех ответов. Преподавателям отправляйте только пустой файл «'+SURVEY+'».')] + notes[5:] + [
            ('Обновление свода','Сохраните полученные анкеты в отдельную папку и загрузите их в Codex либо используйте служебный инструмент. Свод пересобирается при объединении, не в реальном времени.'),
            ('Повторные ответы','Одинаковые копии учитываются один раз. Разные ответы одного человека требуют явного выбора актуальной версии. Имя файла и дата изменения не определяют, какая анкета актуальна.'),
            ('Пустые строки в реестре','Для тем и программ без отметок оставлены строки с пустым ФИО. Так весь перечень остаётся доступен для фильтрации. Пустое ФИО не является отказом преподавателя.'),
            ('Отметка относится к теме','Программы в реестре показывают, где встречается тема. При ограничениях по конкретной программе преподаватель выбирает «Частично» и поясняет их в комментарии. Такие ответы требуют уточнения.'),
            ('Люди с одинаковыми ФИО','Используются только ФИО. Если разные преподаватели имеют полностью одинаковые ФИО, добавьте название организации в скобках в это же поле.'),
        ]
    row=3
    for title,body in notes:
        b.band(s,row,1,7,title,color=b.TEAL,bold=True,height=27)
        b.band(s,row+1,1,7,body,color=b.WHITE,height=50 if len(body)>200 else 42)
        row+=3
    s.freeze_panes='A3';s.print_area=f'A1:G{row}'


def make_survey(data, output):
    wb=new_book()
    s=sheet(wb,'Выбор тем',[15,66,28,20,18,45],title='КАКИЕ ТЕМЫ ВЫ ГОТОВЫ ЧИТАТЬ?',stripe=8)
    last=HEADER_ROW+len(data['topics'])
    b.band(s,2,1,6,'Каталог тем · Rozum · Содержание внутри файла',height=28)
    b.band(s,3,1,6,'Укажите ФИО, откройте содержание и отметьте готовность. Заполняйте светло-зелёные ячейки.',height=29)
    for cell,label in [('A4','Ваше ФИО'),('E4','Дата анкеты'),('A5','Заполнение'),('E5','Версия')]:
        b.text(s[cell],label);s[cell].font=Font(name='Calibri',size=11,bold=True,color=b.INK)
        s[cell].alignment=Alignment(vertical='center',wrap_text=True)
    for ref in ['B4:D4','B5:D5']:s.merge_cells(ref)
    for cell in ['B4','F4','B5']:b.input_cell(s[cell])
    s['F4'].number_format='DD.MM.YYYY'
    b.text(s['F5'],data['version']);b.text(s['B5'],'В процессе')
    for r in [4,5]:s.row_dimensions[r].height=28
    b.band(s,6,1,6,'«Да» — готовы читать тему. При ограничениях по содержанию или программе выберите «Частично» и поясните их в комментарии.',color=b.INPUT,height=31)
    for start,end,label,target in [(1,3,'Все содержания →','Содержание тем'),(4,6,'Инструкция →','Инструкция')]:
        s.merge_cells(start_row=7,start_column=start,end_row=7,end_column=end)
        b.link(s.cell(7,start),label,target,1)
    s.row_dimensions[7].height=27
    rows=[[t['id'],t['title'],'; '.join(codes(t)),'Открыть →','',''] for t in data['topics']]
    b.table(s,HEADERS,rows,start=HEADER_ROW,name='SurveyTopics',heights=60)
    for row,t in enumerate(data['topics'],FIRST_TOPIC_ROW):
        for col in [5,6]:b.input_cell(s.cell(row,col))
        lines=max(len(t['title'])//56+1,len('; '.join(codes(t)))//23+1)
        s.row_dimensions[row].height=max(48,16*lines+16)
    dropdown(s,f'E{FIRST_TOPIC_ROW}:E{last}',['Да','Частично'])
    dropdown(s,'B5',['В процессе','Выбор завершён'])
    s.conditional_formatting.add(f'E{FIRST_TOPIC_ROW}:E{last}',FormulaRule(formula=[f'E{FIRST_TOPIC_ROW}="Да"'],fill=PatternFill('solid',fgColor=b.GREEN)))
    s.conditional_formatting.add(f'F{FIRST_TOPIC_ROW}:F{last}',FormulaRule(formula=[f'AND(E{FIRST_TOPIC_ROW}="Частично",F{FIRST_TOPIC_ROW}="")'],fill=PatternFill('solid',fgColor='C9EBDF')))
    for name,cell in [('LecturerName','B4'),('ResponseDate','F4'),('SurveyVersion','F5'),('CompletionState','B5')]:named(wb,name,cell)
    b.band(s,last+2,1,6,'Сохраните файл «Анкета_Фамилия.xlsx» и отправьте организатору ЛИЧНО. Других ответов преподавателей в этом файле нет.',height=30)
    s.freeze_panes=f'C{FIRST_TOPIC_ROW}';s.print_title_rows=f'1:{HEADER_ROW}';s.print_area=f'A1:F{last+2}'
    anchors=add_details(wb,data,'Выбор тем','E',FIRST_TOPIC_ROW,'A')
    for row,t in enumerate(data['topics'],FIRST_TOPIC_ROW):
        label='Описание не найдено →' if not any(u['content'] for u in t['usages']) else 'Открыть →'
        b.link(s.cell(row,4),label,'Содержание тем',anchors[t['id']])
    add_instructions(wb,data)
    wb.properties.title='Каталог тем Rozum — анкета преподавателя'
    wb.save(output)


def scalar(wb,name):
    if name not in wb.defined_names:raise ValueError(f'Отсутствует поле {name}; используйте полный исходный шаблон.')
    refs=list(wb.defined_names[name].destinations)
    if len(refs)!=1:raise ValueError(f'Неверное поле {name}.')
    sh,ref=refs[0];c=wb[sh][ref]
    if c.data_type=='f':raise ValueError(f'В поле {name} формула вместо значения.')
    return c.value


def read_response(path,data):
    wb=load_workbook(path,data_only=False,read_only=True)
    try:
        if scalar(wb,'SurveyVersion')!=data['version']:raise ValueError('Другая версия каталога. Требуется отдельное сопоставление.')
        name=str(scalar(wb,'LecturerName') or '').strip()
        if not name:raise ValueError('Не заполнено ФИО.')
        state=scalar(wb,'CompletionState')
        if state not in ['В процессе','Выбор завершён']:raise ValueError('Неизвестное состояние заполнения.')
        when=scalar(wb,'ResponseDate')
        if isinstance(when,datetime):when=when.date()
        if isinstance(when,date):when=when.isoformat()
        elif when:
            parsed=None
            for fmt in ['%d.%m.%Y','%Y-%m-%d']:
                try:parsed=datetime.strptime(str(when).strip(),fmt).date().isoformat();break
                except ValueError:pass
            if not parsed:raise ValueError('Дата должна быть датой Excel или текстом ДД.ММ.ГГГГ.')
            when=parsed
        else:when=''
        s=wb['Выбор тем']
        header=next((r[0].row for r in s.iter_rows() if [c.value for c in r[:len(HEADERS)]]==HEADERS),None)
        if not header:raise ValueError('Не найдены заголовки анкеты.')
        topics={t['id']:t for t in data['topics']};seen=set();answers=[]
        for row_number,row in enumerate(s.iter_rows(min_row=header+1,max_col=len(HEADERS)),header+1):
            code=row[0].value
            if code not in topics:
                if any(c.value not in (None,'') for c in [row[4],row[5]]) or str(code or '').startswith('T-'):
                    raise ValueError(f'Неизвестная тема в строке {row_number}.')
                continue
            if code in seen:raise ValueError(f'Повтор кода темы {code}.')
            seen.add(code);t=topics[code]
            if norm(row[1].value)!=norm(t['title']):raise ValueError(f'Изменено название темы {code}.')
            if any(c.data_type=='f' for c in [row[0],row[1],row[4],row[5]]):raise ValueError(f'Формула вместо ответа в теме {code}.')
            ready=norm(row[4].value);comment=str(row[5].value or '').strip()
            if ready not in ['','да','частично']:raise ValueError(f'Неизвестная отметка в теме {code}.')
            if not ready:
                if comment:raise ValueError(f'Комментарий без отметки в теме {code}.')
                continue
            answers.append({'id':code,'ready':'Да' if ready=='да' else 'Частично','programs':codes(t),'comment':comment})
        if seen!=set(topics):raise ValueError('Удалены строки каталога; пришлите полную анкету.')
        return {'name':name,'date':when,'state':state,'answers':sorted(answers,key=lambda a:a['id']),'file':Path(path).name}
    finally:wb.close()


def responses(data,files):
    people={};duplicates=[]
    for path in sorted(files):
        try:r=read_response(path,data)
        except (ValueError,KeyError) as e:raise ValueError(f'{Path(path).name}: {e}') from e
        identity=norm(r['name'])
        nk=norm(r['name'])
        if identity in people:
            p=people[identity]
            if any(p[k]!=r[k] for k in ['answers','date','state']) or norm(p['name'])!=nk:
                raise ValueError(f'Разные версии ответа одного человека: {p["file"]} и {r["file"]}. Оставьте актуальную явно.')
            duplicates.append(r['file'])
        else:people[identity]=r
    return list(people.values()),duplicates


def make_registry(data,files,output):
    people,dups=responses(data,files)
    by_pair=defaultdict(list)
    for p in people:
        for a in p['answers']:
            for program in a['programs']:by_pair[(a['id'],program)].append((p,a))
    wb=new_book();s=sheet(wb,'Реестр',[64,20,35,17,43,80,23,15,17,36],title='ТЕМЫ · ПРОГРАММЫ · ЛЕКТОРЫ')
    b.band(s,2,1,10,f'Получено анкет: {len(people)} · Повторных файлов пропущено: {len(dups)}',height=28)
    b.band(s,3,1,10,'Отметка относится к теме. Программы показывают, где она встречается. Частичная готовность требует уточнения по комментарию.',height=29)
    rows=[];programs={p['code']:p for p in data['programs']}
    for t in data['topics']:
        for program in codes(t):
            found=by_pair[(t['id'],program)]
            for p,a in found:
                rows.append([t['title'],program,p['name'],a['ready'],a['comment'],programs[program]['title'],
                             'Уточнить объём' if a['ready']=='Частично' else '',t['id'],p['date'],p['file']])
            if not found:rows.append([t['title'],program,'','','',programs[program]['title'],'Нет отметок',t['id'],'',''])
    b.table(s,['Тема','Программа','ФИО лектора','Готовность','Комментарий','Полное название программы','Статус','Код темы','Дата анкеты','Файл ответа'],rows,name='Registry',heights=60)
    s.freeze_panes='C6'
    cat=sheet(wb,'Темы',[15,65,33,19,19,20],title='КАТАЛОГ ТЕМ БЕЗ ПОВТОРОВ')
    b.band(cat,2,1,6,'Связи с разными программами и ответы преподавателей находятся на листе «Реестр».',height=28)
    cr=[]
    for t in data['topics']:
        marks=[a for p in people for a in p['answers'] if a['id']==t['id']]
        cr.append([t['id'],t['title'],'; '.join(codes(t)),sum(a['ready']=='Да' for a in marks),sum(a['ready']=='Частично' for a in marks),'Открыть →'])
    b.table(cat,['Код темы','Тема','Программы','Лекторов: да','Частично','Содержание'],cr,name='Topics',heights=60)
    for row,t in enumerate(data['topics'],6):
        lines=max(len(t['title'])//56+1,len('; '.join(codes(t)))//27+1)
        cat.row_dimensions[row].height=max(60,16*lines+16)
    anchors=add_details(wb,data,'Темы','B',6,'A')
    for i,t in enumerate(data['topics'],6):b.link(cat.cell(i,6),'Открыть →','Содержание тем',anchors[t['id']])
    forms=sheet(wb,'Полученные анкеты',[50,20,30,48],title='ПОЛУЧЕННЫЕ АНКЕТЫ')
    b.band(forms,2,1,4,'Здесь показаны только полученные файлы. Кто не прислал ответ, без отдельного списка приглашённых определить нельзя.',height=30)
    b.table(forms,['ФИО','Дата','Заполнение','Файл'],[[p['name'],p['date'],p['state'],p['file']] for p in people],name='Responses',heights=48)
    add_instructions(wb,data,admin=True)
    wb.properties.title='Реестр лекторов Rozum'
    wb.save(output)
    return {'forms':len(people),'marks':sum(len(p['answers']) for p in people),'topic_program_lecturer_links':sum(len(v) for v in by_pair.values()),'duplicates':len(dups)}


def make_audit(data,output):
    wb=new_book()
    s=sheet(wb,'Расхождения',[18,65,30,85,80],title='ПРОВЕРКА ПРОГРАММ И УТП')
    b.band(s,2,1,5,'Особые случаи сопоставления: отсутствующие описания, различия наименований, общие практикумы и темы, добавленные из программы.',height=31)
    b.table(s,['Программа','Тема УТП','Результат сопоставления','Источник в программе','Пояснение'],
            [[u['program'],u['utp_title'],u['status'],' / '.join(u['source_headings']),u['note']] for u in data['exceptions']],name='Issues',heights=90)
    s=sheet(wb,'Все связи',[15,63,17,60,30,80,65],title='ВСЕ СВЯЗИ ТЕМ С ИСТОЧНИКАМИ')
    rows=[]
    for t in data['topics']:
        for u in t['usages']:
            location=(f"{u['section']}; таблица {u['utp_table']}, строка {u['utp_row']}" if u.get('utp_table') is not None else f"{u['section']}; программа, абзацы {', '.join(map(str,u['source_paragraphs']))}")
            rows.append([t['id'],t['title'],u['program'],u['utp_title'],u['status'],' / '.join(u['source_headings']),location])
    b.table(s,['Код темы','Тема каталога','Программа','Название в УТП','Сопоставление','Заголовок в программе','Расположение в источнике'],rows,name='Mappings',heights=64)
    s=sheet(wb,'Объединённые названия',[80,80,60],title='ВАРИАНТЫ НАЗВАНИЙ')
    b.band(s,2,1,3,'Дополнительно нормализованы регистр, пробелы, пунктуация и ё/е. Содержание по программам сохранено отдельно.',height=29)
    b.table(s,['Исходный вариант','Название каталога','Правило'],[[k,v,'Смысловое соответствие; тексты программ сохранены'] for k,v in data['title_aliases'].items()],name='Aliases',heights=64)
    s=sheet(wb,'Похожие, оставлены отдельно',[85,85,55],title='ПОХОЖИЕ ТЕМЫ, ОСТАВЛЕННЫЕ ОТДЕЛЬНО')
    b.table(s,['Тема 1','Тема 2','Причина'],[[p['title1'],p['title2'],'Темы сохранены отдельно из-за различий в содержании или объёме'] for p in data['possible_similar']],name='Similar',heights=80)
    s=sheet(wb,'Документы',[16,44,50,78,20],title='ПРОВЕРЕННЫЕ ДОКУМЕНТЫ')
    rows=[]
    for p in data['programs']:
        for kind in ['program','utp']:rows.append([p['code'],'Программа' if kind=='program' else 'УТП',p[kind+'_file'],p[kind+'_url'],p[kind+'_sha256']])
    b.table(s,['Код','Тип','Имя файла','Ссылка Google Drive','SHA-256'],rows,name='Sources',heights=48)
    wb.save(output)


def main():
    parser=argparse.ArgumentParser(description='Каталог тем Rozum: анкета и объединение ответов.')
    subs=parser.add_subparsers(dest='command',required=True)
    subs.add_parser('build',help='Пересоздать пустой шаблон, пустой реестр и отчёт источников в этой папке')
    merge=subs.add_parser('merge',help='Объединить заполненные анкеты')
    merge.add_argument('folder',type=Path)
    merge.add_argument('--output',type=Path,required=True)
    args=parser.parse_args();data=catalog()
    if args.command=='build':
        make_survey(data,ROOT/SURVEY);make_registry(data,[],ROOT/REGISTRY);make_audit(data,ROOT/AUDIT)
        print(json.dumps(data['stats'],ensure_ascii=False));return
    if args.output.exists():parser.error('Выходной файл существует. Укажите новое имя, чтобы сохранить предыдущий свод.')
    if not args.folder.is_dir():parser.error('Папка с ответами не найдена.')
    files=[p for p in args.folder.glob('*.xlsx') if not p.name.startswith('~$') and p.resolve()!=args.output.resolve()]
    if not files:parser.error('В папке нет анкет .xlsx.')
    try:result=make_registry(data,files,args.output)
    except ValueError as e:parser.exit(1,str(e)+'\n')
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':main()
