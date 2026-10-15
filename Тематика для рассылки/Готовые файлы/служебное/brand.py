"""Create a small, offline lecturer survey and consolidate returned copies."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime
import argparse
import json
from pathlib import Path
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.table import TableStyle, TableStyleElement
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter
from openpyxl.writer.theme import theme_xml
from xml.etree import ElementTree as ET

INK = '221F1F'
TEAL = '008685'
TURQUOISE = '26BBA8'
BRAND_GREEN = '40B870'
PALE = 'F0F8F4'
INPUT = 'E5F5EF'
GRAY = '60756C'
WHITE = 'FFFFFF'
NOTICE = 'F2F5F3'
GREEN = 'DDF2E5'
VERSION = 'DEMO-2026-10-v1'
SURVEY_HEADERS = ['Код темы', 'Тема', 'Программы', 'Содержание', 'Могу читать', 'Комментарий преподавателя']
ROOT = Path(__file__).resolve().parents[1]


def brand_workbook(wb):
    """Set the actual Excel table style and theme, including filter headers."""
    specs = [
        ('wholeTable', DifferentialStyle(font=Font(name='Calibri', size=11, color=INK),
                                         fill=PatternFill('solid', fgColor=WHITE))),
        ('headerRow', DifferentialStyle(font=Font(name='Calibri', size=11, color=WHITE, bold=True),
                                        fill=PatternFill('solid', fgColor=INK),
                                        border=Border(bottom=Side(style='medium', color=TURQUOISE)))),
        ('firstRowStripe', DifferentialStyle(fill=PatternFill('solid', fgColor=WHITE))),
        ('secondRowStripe', DifferentialStyle(fill=PatternFill('solid', fgColor=PALE))),
    ]
    elements = []
    for kind, style in specs:
        elements.append(TableStyleElement(type=kind, dxfId=wb._differential_styles.add(style),
                                         size=1 if kind.endswith('Stripe') else None))
    wb._table_styles.tableStyle.append(TableStyle(name='RozumTable', pivot=False, table=True,
                                                 tableStyleElement=elements))
    wb._table_styles.defaultTableStyle = 'RozumTable'
    ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    theme = ET.fromstring(wb.loaded_theme or theme_xml)
    theme.set('name', 'Rozum')
    scheme = theme.find(f'.//{{{ns}}}clrScheme')
    scheme.set('name', 'Rozum')
    colors = {'dk1': INK, 'lt1': WHITE, 'dk2': INK, 'lt2': PALE,
              'accent1': TURQUOISE, 'accent2': BRAND_GREEN, 'accent3': TEAL,
              'accent4': '009763', 'accent5': GRAY, 'accent6': GREEN,
              'hlink': TEAL, 'folHlink': '009763'}
    for name, color in colors.items():
        target = scheme.find(f'{{{ns}}}{name}')
        target.clear()
        ET.SubElement(target, f'{{{ns}}}srgbClr', {'val': color})
    wb.loaded_theme = ET.tostring(theme, encoding='utf-8', xml_declaration=True)


def brand_stripe(ws, row, columns):
    for col in range(1, columns + 1):
        ws.cell(row, col).fill = PatternFill('solid', fgColor=TURQUOISE if col <= (columns+1)//2 else BRAND_GREEN)
    ws.row_dimensions[row].height = 6


def brand_title(ws, columns, title):
    band(ws, 1, 1, columns - 1, title, color=WHITE, size=20, height=52, bold=True)
    logo = ROOT / 'служебное' / 'оформление' / 'rozum_logo.png'
    if logo.exists():
        im = Image(str(logo))
        width = 170
        height = round(width * im.height / im.width)
        im.anchor = OneCellAnchor(_from=AnchorMarker(col=columns-1, row=0, colOff=6*9525, rowOff=15*9525),
                                 ext=XDRPositiveSize2D(cx=width*9525, cy=height*9525))
        ws.add_image(im)


def text(cell, value):
    cell.value = '' if value is None else str(value)
    cell.data_type = 's'


def base_sheet(ws, widths):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.35
    ws.oddFooter.center.text = 'rozum · Пробная версия · &P / &N'
    ws.sheet_properties.outlinePr.summaryRight = False


def band(ws, row, start, end, value, *, color=PALE, size=11, height=30, bold=False):
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    for cells in ws.iter_rows(min_row=row, max_row=row, min_col=start, max_col=end):
        for c in cells:
            c.fill = PatternFill('solid', fgColor=color)
    c = ws.cell(row, start)
    text(c, value)
    c.font = Font(name='Calibri', size=size, bold=bold, color=WHITE if color in (INK, TEAL) else INK)
    c.alignment = Alignment(vertical='center', wrap_text=True, indent=1)
    ws.row_dimensions[row].height = height


def link(cell, label, sheet, row, column='A'):
    text(cell, label)
    cell.hyperlink = None
    from openpyxl.worksheet.hyperlink import Hyperlink
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{column}{row}", display=label)
    cell.font = Font(name='Calibri', size=11, color=TEAL, bold=True, underline='single')
    cell.alignment = Alignment(wrap_text=True, vertical='center')


def input_cell(cell):
    cell.fill = PatternFill('solid', fgColor=INPUT)
    cell.font = Font(name='Calibri', size=11, color=INK)
    cell.alignment = Alignment(vertical='center', wrap_text=True)
    cell.border = Border(bottom=Side(style='thin', color=BRAND_GREEN))


def named(wb, name, target):
    wb.defined_names.add(DefinedName(name, attr_text=target))


def table(ws, headers, rows, *, start=5, name='DataTable', heights=54):
    for j, h in enumerate(headers, 1):
        c = ws.cell(start, j)
        text(c, h)
        c.fill = PatternFill('solid', fgColor=INK)
        c.font = Font(name='Calibri', size=11, color=WHITE, bold=True)
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[start].height = 32
    for i, values in enumerate(rows, start + 1):
        ws.row_dimensions[i].height = heights
        for j, value in enumerate(values, 1):
            c = ws.cell(i, j)
            if isinstance(value, (int, float, date, datetime)):
                c.value = value
            else:
                text(c, value)
            c.font = Font(name='Calibri', size=11, color=INK)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.fill = PatternFill('solid', fgColor=WHITE if i % 2 else PALE)
    if rows:
        t = Table(displayName=name, ref=f'A{start}:{get_column_letter(len(headers))}{start+len(rows)}')
        t.tableStyleInfo = TableStyleInfo(name='RozumTable', showRowStripes=False)
        ws.add_table(t)
    else:
        ws.auto_filter.ref = f'A{start}:{get_column_letter(len(headers))}{start}'
    ws.freeze_panes = f'C{start+1}'
    ws.print_title_rows = f'1:{start}'
    ws.print_area = f'A1:{get_column_letter(len(headers))}{max(start+len(rows),start+1)}'

