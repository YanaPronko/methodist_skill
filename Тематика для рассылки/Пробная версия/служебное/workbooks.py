"""Create a small, offline lecturer survey and consolidate returned copies."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime
import argparse
import json
from pathlib import Path
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.table import TableStyle, TableStyleElement
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter
from openpyxl.writer.theme import theme_xml
from xml.etree import ElementTree as ET

INK = '221F1F'
TEAL = '008685'
TURQUOISE = '26BBA8'
BRAND_GREEN = '40B870'
PALE = 'F0F8F4'
INPUT = 'E5F5EF'
GRAY = '60756C'
WHITE = 'FFFFFF'
NOTICE = 'F2F5F3'
GREEN = 'DDF2E5'
VERSION = 'DEMO-2026-10-v1'
SURVEY_HEADERS = ['Код темы', 'Тема', 'Программы', 'Содержание', 'Могу читать', 'Комментарий преподавателя']
ROOT = Path(__file__).resolve().parents[1]


def brand_workbook(wb):
    """Set the actual Excel table style and theme, including filter headers."""
    specs = [
        ('wholeTable', DifferentialStyle(font=Font(name='Calibri', size=11, color=INK),
                                         fill=PatternFill('solid', fgColor=WHITE))),
        ('headerRow', DifferentialStyle(font=Font(name='Calibri', size=11, color=WHITE, bold=True),
                                        fill=PatternFill('solid', fgColor=INK),
                                        border=Border(bottom=Side(style='medium', color=TURQUOISE)))),
        ('firstRowStripe', DifferentialStyle(fill=PatternFill('solid', fgColor=WHITE))),
        ('secondRowStripe', DifferentialStyle(fill=PatternFill('solid', fgColor=PALE))),
    ]
    elements = []
    for kind, style in specs:
        elements.append(TableStyleElement(type=kind, dxfId=wb._differential_styles.add(style),
                                         size=1 if kind.endswith('Stripe') else None))
    wb._table_styles.tableStyle.append(TableStyle(name='RozumTable', pivot=False, table=True,
                                                 tableStyleElement=elements))
    wb._table_styles.defaultTableStyle = 'RozumTable'
    ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    theme = ET.fromstring(wb.loaded_theme or theme_xml)
    theme.set('name', 'Rozum')
    scheme = theme.find(f'.//{{{ns}}}clrScheme')
    scheme.set('name', 'Rozum')
    colors = {'dk1': INK, 'lt1': WHITE, 'dk2': INK, 'lt2': PALE,
              'accent1': TURQUOISE, 'accent2': BRAND_GREEN, 'accent3': TEAL,
              'accent4': '009763', 'accent5': GRAY, 'accent6': GREEN,
              'hlink': TEAL, 'folHlink': '009763'}
    for name, color in colors.items():
        target = scheme.find(f'{{{ns}}}{name}')
        target.clear()
        ET.SubElement(target, f'{{{ns}}}srgbClr', {'val': color})
    wb.loaded_theme = ET.tostring(theme, encoding='utf-8', xml_declaration=True)


def brand_stripe(ws, row, columns):
    for col in range(1, columns + 1):
        ws.cell(row, col).fill = PatternFill('solid', fgColor=TURQUOISE if col <= (columns+1)//2 else BRAND_GREEN)
    ws.row_dimensions[row].height = 6


def brand_title(ws, columns, title):
    band(ws, 1, 1, columns - 1, title, color=WHITE, size=20, height=52, bold=True)
    logo = ROOT / 'служебное' / 'оформление' / 'rozum_logo.png'
    if logo.exists():
        im = Image(logo)
        width = 170
        height = round(width * im.height / im.width)
        im.anchor = OneCellAnchor(_from=AnchorMarker(col=columns-1, row=0, colOff=6*9525, rowOff=15*9525),
                                 ext=XDRPositiveSize2D(cx=width*9525, cy=height*9525))
        ws.add_image(im)


def text(cell, value):
    cell.value = '' if value is None else str(value)
    cell.data_type = 's'


def base_sheet(ws, widths):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.35
    ws.oddFooter.center.text = 'rozum · Пробная версия · &P / &N'
    ws.sheet_properties.outlinePr.summaryRight = False


def band(ws, row, start, end, value, *, color=PALE, size=11, height=30, bold=False):
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    for cells in ws.iter_rows(min_row=row, max_row=row, min_col=start, max_col=end):
        for c in cells:
            c.fill = PatternFill('solid', fgColor=color)
    c = ws.cell(row, start)
    text(c, value)
    c.font = Font(name='Calibri', size=size, bold=bold, color=WHITE if color in (INK, TEAL) else INK)
    c.alignment = Alignment(vertical='center', wrap_text=True, indent=1)
    ws.row_dimensions[row].height = height


def link(cell, label, sheet, row, column='A'):
    text(cell, label)
    cell.hyperlink = None
    from openpyxl.worksheet.hyperlink import Hyperlink
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{column}{row}", display=label)
    cell.font = Font(name='Calibri', size=11, color=TEAL, bold=True, underline='single')
    cell.alignment = Alignment(wrap_text=True, vertical='center')


def input_cell(cell):
    cell.fill = PatternFill('solid', fgColor=INPUT)
    cell.font = Font(name='Calibri', size=11, color=INK)
    cell.alignment = Alignment(vertical='center', wrap_text=True)
    cell.border = Border(bottom=Side(style='thin', color=BRAND_GREEN))


def named(wb, name, target):
    wb.defined_names.add(DefinedName(name, attr_text=target))


def table(ws, headers, rows, *, start=5, name='DataTable', heights=54):
    for j, h in enumerate(headers, 1):
        c = ws.cell(start, j)
        text(c, h)
        c.fill = PatternFill('solid', fgColor=INK)
        c.font = Font(name='Calibri', size=11, color=WHITE, bold=True)
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[start].height = 32
    for i, values in enumerate(rows, start + 1):
        ws.row_dimensions[i].height = heights
        for j, value in enumerate(values, 1):
            c = ws.cell(i, j)
            if isinstance(value, (int, float, date, datetime)):
                c.value = value
            else:
                text(c, value)
            c.font = Font(name='Calibri', size=11, color=INK)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.fill = PatternFill('solid', fgColor=WHITE if i % 2 else PALE)
    if rows:
        t = Table(displayName=name, ref=f'A{start}:{get_column_letter(len(headers))}{start+len(rows)}')
        t.tableStyleInfo = TableStyleInfo(name='RozumTable', showRowStripes=False)
        ws.add_table(t)
    else:
        ws.auto_filter.ref = f'A{start}:{get_column_letter(len(headers))}{start}'
    ws.freeze_panes = f'C{start+1}'
    ws.print_title_rows = f'1:{start}'
    ws.print_area = f'A1:{get_column_letter(len(headers))}{max(start+len(rows),start+1)}'


def survey(catalog, output):
    wb = Workbook()
    brand_workbook(wb)
    ws = wb.active
    ws.title = 'Выбор тем'
    details = wb.create_sheet('Содержание тем')
    how = wb.create_sheet('Инструкция')
    base_sheet(ws, [15, 66, 22, 18, 18, 42])
    base_sheet(details, [15, 29, 29, 29, 29, 26])
    base_sheet(how, [24, 30, 30, 30, 30, 24])
    ws.sheet_properties.tabColor = TEAL
    details.sheet_properties.tabColor = INK
    how.sheet_properties.tabColor = GRAY
    brand_title(ws, 6, 'КАКИЕ ТЕМЫ ВЫ ГОТОВЫ ЧИТАТЬ?')
    band(ws, 2, 1, 6, f"Пробная анкета · {len(catalog['topics'])} тем · 1 программа · работает без интернета", height=28)
    band(ws, 3, 1, 6, 'Укажите ФИО → откройте содержание → выберите «Да» или «Частично» → сохраните свою копию.', height=29)
    for cell, value in [('A4', 'Ваше ФИО'), ('E4', 'Дата анкеты'), ('A5', 'Версия'), ('E5', 'Заполнение')]:
        text(ws[cell], value)
        ws[cell].font = Font(name='Calibri', size=11, bold=True, color=INK)
        ws[cell].alignment = Alignment(wrap_text=True, vertical='center')
    ws.merge_cells('B4:D4')
    ws.merge_cells('B5:D5')
    input_cell(ws['B4']); input_cell(ws['F4']); input_cell(ws['F5'])
    ws['F4'].number_format = 'DD.MM.YYYY'
    text(ws['F5'], 'В процессе')
    text(ws['B5'], VERSION)
    ws.row_dimensions[4].height = ws.row_dimensions[5].height = 30
    band(ws, 6, 1, 6, 'Светло-зелёные ячейки — для ваших ответов. Пустая отметка означает «не указано». При ответе «Частично» поясните ограничения.', color=INPUT, height=29)
    band(ws, 7, 1, 6, 'ОБРАЗЕЦ: темы из локальной программы ПТИБ (advance), не из папок Google Drive. Для проверки формата; не полная рассылка.', color=NOTICE, height=29)
    text(ws['A8'], 'Выбрано тем:')
    ws['B8'] = '=COUNTIF(E11:E22,"Да")+COUNTIF(E11:E22,"Частично")'
    ws['B8'].font = Font(name='Calibri', size=12, bold=True, color=TEAL)
    link(ws['D8'], 'Как заполнить →', 'Инструкция', 1)
    link(ws['F8'], 'Все содержания →', 'Содержание тем', 1)
    rows = [[t['id'], t['title'], catalog['program']['short'], 'Открыть →', '', ''] for t in catalog['topics']]
    table(ws, SURVEY_HEADERS, rows, start=10, name='SurveyTopics', heights=68)
    for i, topic in enumerate(catalog['topics'], 11):
        input_cell(ws.cell(i, 5)); input_cell(ws.cell(i, 6))
        ws.row_dimensions[i].height = max(60, (len(topic['title']) // 57 + 1) * 17 + 16)
    band(ws, 24, 1, 6, 'Сохраните файл как «Анкета_Фамилия.xlsx» и отправьте организатору ЛИЧНО. Другие преподаватели ваши ответы не увидят.', height=31)
    ws.print_area = 'A1:F24'
    ws.page_setup.fitToHeight = 1
    ws.print_title_rows = '1:10'
    ws.freeze_panes = 'C11'
    brand_stripe(ws, 9, 6)
    ready = DataValidation(type='list', formula1='"Да,Частично"', allow_blank=True)
    ready.errorTitle = 'Выберите значение из списка'
    ready.error = 'Допустимы «Да», «Частично» или пустая ячейка.'
    ready.showErrorMessage = True
    ready.showInputMessage = False
    ready.errorStyle = 'stop'
    ready.showDropDown = False
    ws.add_data_validation(ready); ready.add('E11:E22')
    done = DataValidation(type='list', formula1='"В процессе,Выбор завершён"', allow_blank=False)
    done.showErrorMessage = True; done.errorStyle = 'stop'; done.showDropDown = False
    done.showInputMessage = False
    ws.add_data_validation(done); done.add('F5')
    ws.conditional_formatting.add('E11:E22', FormulaRule(formula=['E11="Да"'], fill=PatternFill('solid', fgColor=GREEN)))
    ws.conditional_formatting.add('F11:F22', FormulaRule(formula=['AND(E11="Частично",F11="")'], fill=PatternFill('solid', fgColor='C9EBDF')))
    for name, ref in [('LecturerName', '$B$4'), ('ResponseDate', '$F$4'), ('SurveyVersion', '$B$5'), ('CompletionState', '$F$5')]:
        named(wb, name, f"'Выбор тем'!{ref}")
    brand_title(details, 6, 'СОДЕРЖАНИЕ ТЕМ')
    brand_stripe(details, 4, 6)
    band(details, 2, 1, 6, 'Текст перенесён из программы-образца. Ниже каждая тема показана вместе с её контекстом.', height=28)
    link(details['A3'], '← К выбору тем', 'Выбор тем', 1)
    details.row_dimensions[3].height = 24
    for index, topic in enumerate(catalog['topics']):
        row = 5 + index * 7
        band(details, row, 1, 6, topic['title'], color=TEAL, size=13, height=49, bold=True)
        band(details, row + 1, 1, 6, f"{topic['id']} · {topic['section']} · {topic['original_heading']}", height=34)
        band(details, row + 2, 1, 6, 'Программа: ' + catalog['program']['title'], height=38)
        band(details, row + 3, 1, 6, topic['content'], color=WHITE, size=12, height=max(76, 20 * (len(topic['content']) // 130 + topic['content'].count('\n') + 1)))
        band(details, row + 4, 1, 6, 'Источник: ' + catalog['source']['file'] + ' · ' + topic['section'] + f" · Тема {topic['number']}", height=28)
        link(details.cell(row + 5, 1), '← Отметить тему', 'Выбор тем', index + 11, 'E')
        back = details.cell(row + 5, 1)
        back.hyperlink = None
        back.value = f'=HYPERLINK("#\'Выбор тем\'!E"&MATCH("{topic["id"]}",\'Выбор тем\'!$A:$A,0),"← Отметить тему")'
        details.merge_cells(start_row=row + 5, start_column=1, end_row=row + 5, end_column=3)
        details.row_dimensions[row + 5].height = 25
        link(ws.cell(index + 11, 4), 'Открыть →', 'Содержание тем', row)
        if index and index % 2 == 0:
            details.row_breaks.append(Break(id=row - 1))
    details.freeze_panes = 'A5'
    details.print_title_rows = '1:3'
    details.print_area = f'A1:F{5+len(catalog["topics"])*7-1}'
    brand_title(how, 6, 'КАК ЗАПОЛНИТЬ АНКЕТУ')
    brand_stripe(how, 2, 6)
    instructions = [
        ('01 · Сохраните копию', 'Скачайте этот файл и откройте в Excel или LibreOffice Calc. В просмотрщике чата заполнение может быть недоступно.'),
        ('02 · Представьтесь', 'На листе «Выбор тем» укажите ФИО полностью. При желании добавьте дату. Заполняйте только светло-зелёные ячейки.'),
        ('03 · Изучите темы', 'Нажмите «Открыть →» напротив темы. Содержание откроется внутри этого файла. Ссылка «← Отметить тему» вернёт к ответу.'),
        ('04 · Сделайте выбор', '«Да» — готовы читать всю тему. «Частично» — готовы читать часть: поясните, какую. Пустая ячейка означает «не указано», это не отказ.'),
        ('05 · Сохраните результат', 'Можно вернуться позже. Когда закончите, выберите «Выбор завершён». Сохраните в формате .xlsx как «Анкета_Фамилия.xlsx».'),
        ('06 · Отправьте лично', 'Передайте заполненный файл организатору лично. В общий чат предназначен только пустой шаблон. В файле нет ответов других преподавателей.'),
        ('О пробной версии', '12 тем из одной локальной программы-образца ПТИБ (advance). Содержание перенесено из источника без смыслового пересказа. Папки Google Drive пока недоступны.'),
        ('Для организатора', 'Сводный реестр — отдельный файл. Здесь нет макросов, внешних подключений и сведений из прежней таблицы «Тема-Лектор_1».'),
    ]
    for i, (label, body) in enumerate(instructions):
        row = 3 + i * 3
        band(how, row, 1, 6, label, color=TEAL, bold=True, height=27)
        band(how, row + 1, 1, 6, body, color=WHITE, height=48)
    link(how['A28'], '← К выбору тем', 'Выбор тем', 1)
    how.print_area = 'A1:F28'
    how.freeze_panes = 'A3'
    wb.properties.title = 'Пробная анкета выбора тем преподавателем'
    wb.properties.creator = 'Каталог тем'
    wb.properties.description = 'Офлайн-анкета. Пробная версия. Исходные содержания — из локальной программы-образца.'
    wb.active = 0
    wb.save(output)


def normalize(value):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', str(value or ''))).strip().casefold()


def scalar(wb, name):
    if name not in wb.defined_names:
        raise ValueError(f'Не найдено поле {name}; используйте исходный шаблон.')
    destinations = list(wb.defined_names[name].destinations)
    if len(destinations) != 1:
        raise ValueError(f'Неверное поле {name}.')
    sheet, cell = destinations[0]
    c = wb[sheet][cell]
    if c.data_type == 'f':
        raise ValueError(f'В поле {name} недопустима формула.')
    return c.value


def read_response(path, catalog):
    wb = load_workbook(path, data_only=False)
    try:
        if scalar(wb, 'SurveyVersion') != catalog['version']:
            raise ValueError('Другая версия каталога; требуется отдельное сопоставление тем.')
        name = str(scalar(wb, 'LecturerName') or '').strip()
        if not name:
            raise ValueError('Не заполнено ФИО.')
        state = scalar(wb, 'CompletionState')
        if state not in ('В процессе', 'Выбор завершён'):
            raise ValueError('Неизвестное состояние заполнения.')
        response_date = scalar(wb, 'ResponseDate')
        if isinstance(response_date, datetime):
            response_date = response_date.date()
        if response_date:
            if isinstance(response_date, date):
                response_date = response_date.isoformat()
            else:
                parsed = None
                for fmt in ('%d.%m.%Y', '%Y-%m-%d'):
                    try:
                        parsed = datetime.strptime(str(response_date).strip(), fmt).date().isoformat()
                        break
                    except ValueError:
                        pass
                if not parsed:
                    raise ValueError('Дата должна быть датой Excel либо текстом ДД.ММ.ГГГГ.')
                response_date = parsed
        else:
            response_date = ''
        ws = wb['Выбор тем']
        header = next((row[0].row for row in ws.iter_rows() if [c.value for c in row[:6]] == SURVEY_HEADERS), None)
        if not header:
            raise ValueError('Не найдены заголовки анкеты.')
        by_id = {t['id']: t for t in catalog['topics']}
        seen, answers = set(), []
        for row in ws.iter_rows(min_row=header + 1, max_col=6):
            code = row[0].value
            if code not in by_id:
                if row[4].value not in (None, '') or row[5].value not in (None, '') or str(code or '').startswith('T-'):
                    raise ValueError(f'Неизвестная тема в строке {row[0].row}: {code}')
                continue
            if code in seen:
                raise ValueError(f'Повтор кода темы {code}.')
            seen.add(code)
            if normalize(row[1].value) != normalize(by_id[code]['title']):
                raise ValueError(f'Изменено название темы {code}; требуется проверка.')
            if any(c.data_type == 'f' for c in (row[0], row[1], row[4], row[5])):
                raise ValueError(f'Формула вместо ответа в теме {code}.')
            value = normalize(row[4].value)
            comment = str(row[5].value or '').strip()
            if value not in ('', 'да', 'частично'):
                raise ValueError(f'Неизвестная отметка «{row[4].value}» в теме {code}.')
            if not value and comment:
                raise ValueError(f'Комментарий без отметки в теме {code}; уточните ответ.')
            if value:
                answers.append({'id': code, 'ready': 'Да' if value == 'да' else 'Частично', 'comment': comment})
        if seen != set(by_id):
            raise ValueError('Удалены строки каталога; пришлите полную анкету.')
        return {'name': name, 'key': normalize(name), 'date': response_date, 'state': state,
                'answers': sorted(answers, key=lambda a: a['id']), 'file': path.name}
    finally:
        wb.close()


def consolidate(catalog, files, output):
    responses, duplicates = {}, []
    if not files:
        raise ValueError('Нет входных анкет .xlsx.')
    for path in sorted(files):
        try:
            response = read_response(path, catalog)
        except (ValueError, KeyError) as exc:
            raise ValueError(f'{path.name}: {exc}') from exc
        previous = responses.get(response['key'])
        if previous:
            fields = ('answers', 'date', 'state')
            if any(previous[k] != response[k] for k in fields):
                raise ValueError(f'Две разные анкеты одного ФИО: {previous["file"]} и {response["file"]}. Оставьте актуальную явно; дата файла не используется для выбора.')
            duplicates.append(path.name)
        else:
            responses[response['key']] = response
    wb = Workbook()
    brand_workbook(wb)
    ws = wb.active; ws.title = 'Сводный реестр'
    cat = wb.create_sheet('Каталог тем')
    forms = wb.create_sheet('Полученные анкеты')
    programs = wb.create_sheet('Программы')
    how = wb.create_sheet('Пояснения')
    base_sheet(ws, [15, 64, 23, 35, 18, 44, 20, 16, 35])
    base_sheet(cat, [15, 62, 24, 82, 15, 18, 26])
    base_sheet(forms, [36, 20, 22, 15, 15, 46])
    base_sheet(programs, [22, 26, 95, 46])
    base_sheet(how, [28, 28, 28, 28, 28, 28])
    for sheet, columns in [(ws, 9), (cat, 7), (forms, 6), (programs, 4), (how, 6)]:
        brand_title(sheet, columns, 'ПРОБНЫЙ СВОДНЫЙ РЕЕСТР' if sheet == ws else sheet.title.upper())
        brand_stripe(sheet, 4, columns)
    by_id = {t['id']: t for t in catalog['topics']}
    merged, counts = [], Counter()
    for response in responses.values():
        for a in response['answers']:
            t = by_id[a['id']]
            counts[(a['id'], a['ready'])] += 1
            merged.append([t['id'], t['title'], catalog['program']['short'], response['name'], a['ready'], a['comment'],
                           'Да: уточнить объём' if a['ready'] == 'Частично' else '', response['date'], response['file']])
    merged.sort(key=lambda r: (normalize(r[1]), normalize(r[3])))
    demo = all(r['name'].startswith('Демо-лектор') for r in responses.values())
    sample_note = 'Лекторы в этом примере вымышлены.' if demo else 'Данные из полученных анкет; прежние назначения не импортированы.'
    band(ws, 2, 1, 9, f'Получено анкет: {len(responses)} · Отметок: {len(merged)} · Повторных файлов пропущено: {len(duplicates)}. {sample_note}', color=NOTICE, height=29)
    band(ws, 3, 1, 9, 'Фильтры в заголовках: тема, программа, ФИО, готовность. «Частично» требует уточнения объёма; пустая анкета не является отказом.', height=29)
    table(ws, ['Код темы', 'Тема', 'Программа', 'ФИО лектора', 'Готовность', 'Комментарий', 'Требует уточнения', 'Дата анкеты', 'Файл ответа'], merged, name='LecturerTopicRegistry', heights=68)
    cat_rows = []
    for t in catalog['topics']:
        yes, part = counts[(t['id'], 'Да')], counts[(t['id'], 'Частично')]
        status = 'Есть подтверждение' if yes else ('Только частично' if part else 'Нет отметок')
        cat_rows.append([t['id'], t['title'], catalog['program']['short'], t['content'], yes, part, status])
    band(cat, 2, 1, 7, 'Одна строка на тему. «Нет отметок» — нет ответа, а не подтверждённое отсутствие компетенции.', height=28)
    table(cat, ['Код темы', 'Тема', 'Программа', 'Содержание', 'Да', 'Частично', 'Покрытие темы'], cat_rows, name='TopicCoverage', heights=96)
    form_rows = [[r['name'], r['date'], r['state'], sum(a['ready'] == 'Да' for a in r['answers']), sum(a['ready'] == 'Частично' for a in r['answers']), r['file']] for r in responses.values()]
    band(forms, 2, 1, 6, 'Показаны только полученные файлы. Кто не прислал анкету, по этим данным определить нельзя.', height=28)
    table(forms, ['ФИО лектора', 'Дата анкеты', 'Заполнение', 'Тем полностью', 'Тем частично', 'Файл ответа'], form_rows, name='ReceivedForms', heights=45)
    band(programs, 2, 1, 4, 'Полные названия программ. В остальных таблицах используется сокращение для удобства просмотра.', height=29)
    table(programs, ['Код программы', 'Сокращение', 'Полное название программы', 'Файл-источник'],
          [[catalog['program']['id'], catalog['program']['short'], catalog['program']['title'], catalog['source']['file']]], name='Programs', heights=72)
    notes = [
        ('Это демонстрация объединения специально заполненных копий. ФИО вымышлены. Прежние назначения из «Тема-Лектор_1» не использованы.' if demo
         else 'Реестр составлен из полученных анкет. Прежние назначения из «Тема-Лектор_1» не использованы.'),
        'Рассылайте только «01_Анкета_лектора_ПРОБНАЯ.xlsx». Этот сводный реестр и заполненные копии предназначены организатору.',
        'Полная готовность «Да» и частичная готовность различаются. При частичной готовности организатор уточняет объём, даже если есть комментарий.',
        'Коды тем постоянны для данного каталога. Порядок строк не влияет на объединение. Названия и содержание берутся из эталонного каталога.',
        'Одинаковые копии одного ФИО учитываются один раз. Если два файла одного ФИО отличаются, объединение останавливается: актуальный ответ нужно выбрать явно.',
        'ФИО — идентификатор только для прототипа. Полных однофамильцев с одинаковыми ФИО перед реальной рассылкой нужно различать дополнительным кодом или контактом.',
        'Для нового свода загрузите полученные анкеты в чат. Для самостоятельного объединения есть служебный скрипт и инструкция рядом с файлами.',
        'Свод является снимком полученных ответов. Изменение анкет не обновляет его автоматически: объединение нужно повторить.',
        'Содержание взято из локального образца, не из Google Drive. Пробный каталог содержит одну программу; проверка межпрограммных совпадений ещё не выполнена.',
    ]
    for i, note in enumerate(notes, 3):
        band(how, i, 1, 6, note, color=WHITE if i % 2 else PALE, height=47)
    how.print_area = f'A1:F{len(notes)+2}'
    wb.properties.creator = 'Каталог тем'
    wb.properties.title = 'Пробное объединение ответов преподавателей'
    wb.save(output)
    return {'forms': len(responses), 'marks': len(merged), 'duplicates': len(duplicates)}


def main():
    parser = argparse.ArgumentParser(description='Объединить копии пробной анкеты преподавателей.')
    parser.add_argument('folder', type=Path, help='Папка с заполненными анкетами .xlsx')
    parser.add_argument('--output', type=Path, required=True, help='Путь к новому сводному файлу .xlsx')
    args = parser.parse_args()
    catalog = json.loads((ROOT / 'служебное' / 'каталог.json').read_text(encoding='utf-8'))
    files = [p for p in args.folder.glob('*.xlsx') if p.resolve() != args.output.resolve() and not p.name.startswith('~$')]
    if args.output.exists():
        parser.error('Выходной файл уже существует. Укажите новое имя, чтобы сохранить предыдущий свод.')
    try:
        result = consolidate(catalog, files, args.output)
    except ValueError as exc:
        parser.exit(1, str(exc) + '\n')
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
